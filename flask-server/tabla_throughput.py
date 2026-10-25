import os
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

base = Path(__file__).parent


def leer(nombre):
    f = base / f"resultados_throughput_{nombre}.csv"
    if not f.exists():
        return {}
    return {int(r["clientes_solicitados"]): r for r in csv.DictReader(f.open(encoding="utf-8"))}


def num(valor):
    try:
        return int(valor)
    except (ValueError, TypeError):
        try:
            return float(valor)
        except (ValueError, TypeError):
            return "n/d"


def estado(r):
    if r is None or int(r["emisores"]) == 0:
        return "SIN MEDIR"
    errores = int(r["errores_rechazo"]) + int(r["errores_timeout"])
    ok = (int(r["conexiones_fallidas"]) == 0 and errores == 0
          and int(r["desconexiones_inesperadas"]) == 0
          and r["mensajes_recibidos"] == r["mensajes_esperados"])
    return "OK" if ok else "CON PROBLEMAS"


datos = {"Flask": leer("flask"), "Node.js": leer("node")}
escenarios = sorted({n for d in datos.values() for n in d})

# ---------- CSV y tabla en pantalla (formato largo) ----------
columnas = ["Clientes", "Servidor", "Conectados", "Fallidas", "Enviadas", "Aceptadas",
            "Errores", "Recibidos", "Esperados", "Desconexiones",
            "Ingesta (alertas/s)", "Entrega (msgs/s)"]
filas = []
for n in escenarios:
    for srv, d in datos.items():
        r = d.get(n)
        if not r:
            filas.append([n, srv, "no medido"] + [""] * 9)
            continue
        err = int(r["errores_rechazo"]) + int(r["errores_timeout"])
        filas.append([n, srv, r["conexiones_simultaneas"], r["conexiones_fallidas"],
                      r["alertas_enviadas"], r["alertas_aceptadas"], err,
                      r["mensajes_recibidos"], r["mensajes_esperados"],
                      r["desconexiones_inesperadas"],
                      r["throughput_ingesta_alertas_s"] or "n/d",
                      r["throughput_entrega_msgs_s"] or "n/d"])

print("| " + " | ".join(columnas) + " |")
print("|" + "---|" * len(columnas))
for f in filas:
    print("| " + " | ".join(str(x) for x in f) + " |")

archivo_csv = base / "tabla_comparativa_throughput.csv"
with archivo_csv.open("w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(columnas)
    w.writerows(filas)

# ---------- Excel ----------
AZUL, VERDE, NARANJA, ROJO = "245A81", "2E7D6B", "B26A00", "8E3B46"
COLOR_FLASK, COLOR_NODE = "D35400", "2874A6"
fino = Side(style="thin", color="BFC9D4")
borde = Border(left=fino, right=fino, top=fino, bottom=fino)
centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

libro = Workbook()
hoja = libro.active
hoja.title = "Comparativo"
hoja.sheet_view.showGridLines = False

metricas = [
    ("CONCURRENCIA", AZUL, "Conexiones\nsimultáneas", lambda r: num(r["conexiones_simultaneas"]), "#,##0"),
    ("INGESTA", NARANJA, "Alertas/s\naceptadas", lambda r: num(r["throughput_ingesta_alertas_s"]), "#,##0.00"),
    ("ENTREGA", NARANJA, "Mensajes/s\nentregados", lambda r: num(r["throughput_entrega_msgs_s"]), "#,##0.00"),
    ("ERRORES", ROJO, "Errores", lambda r: int(r["errores_rechazo"]) + int(r["errores_timeout"]), "0"),
    ("DESCONEXIONES", ROJO, "Inesperadas", lambda r: num(r["desconexiones_inesperadas"]), "0"),
]
total_cols = 1 + 2 * len(metricas) + 2

hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
hoja["A1"] = "Comparativo de throughput y concurrencia — Flask vs Node.js"
hoja["A1"].font = Font(bold=True, size=16, color="FFFFFF")
hoja["A1"].fill = PatternFill("solid", fgColor=AZUL)
hoja["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
hoja.row_dimensions[1].height = 32

hoja.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols) 
hoja["A2"].font = Font(italic=True, size=10, color="555555")
hoja["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

# Encabezados: fila 4 grupos, fila 5 servidor
c = hoja.cell(row=4, column=1, value="CLIENTES")
hoja.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)
c.font = Font(bold=True, color="FFFFFF")
c.alignment = centro
for fila_h in (4, 5):
    hoja.cell(row=fila_h, column=1).fill = PatternFill("solid", fgColor=AZUL)
    hoja.cell(row=fila_h, column=1).border = borde

col = 2
for grupo, color, subtitulo, _, _ in metricas:
    hoja.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
    g = hoja.cell(row=4, column=col, value=f"{grupo}\n({subtitulo.replace(chr(10), ' ')})")
    g.font = Font(bold=True, color="FFFFFF", size=10)
    g.alignment = centro
    for k, (srv, colr) in enumerate((("Flask", COLOR_FLASK), ("Node.js", COLOR_NODE))):
        hoja.cell(row=4, column=col + k).fill = PatternFill("solid", fgColor=color)
        hoja.cell(row=4, column=col + k).border = borde
        s = hoja.cell(row=5, column=col + k, value=srv)
        s.font = Font(bold=True, color="FFFFFF")
        s.fill = PatternFill("solid", fgColor=colr)
        s.alignment = centro
        s.border = borde
    col += 2
hoja.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
g = hoja.cell(row=4, column=col, value="ESTADO")
g.font = Font(bold=True, color="FFFFFF", size=10)
g.alignment = centro
for k, (srv, colr) in enumerate((("Flask", COLOR_FLASK), ("Node.js", COLOR_NODE))):
    hoja.cell(row=4, column=col + k).fill = PatternFill("solid", fgColor="4A4A4A")
    hoja.cell(row=4, column=col + k).border = borde
    s = hoja.cell(row=5, column=col + k, value=srv)
    s.font = Font(bold=True, color="FFFFFF")
    s.fill = PatternFill("solid", fgColor=colr)
    s.alignment = centro
    s.border = borde
hoja.row_dimensions[4].height = 36
hoja.row_dimensions[5].height = 22

# Datos
fila_datos = 6
for n in escenarios:
    hoja.cell(row=fila_datos, column=1, value=n)
    col = 2
    for _, _, _, getter, fmt in metricas:
        for srv in ("Flask", "Node.js"):
            r = datos[srv].get(n)
            valor = getter(r) if r else "n/d"
            hoja.cell(row=fila_datos, column=col, value=valor).number_format = fmt
            col += 1
    for srv in ("Flask", "Node.js"):
        est = estado(datos[srv].get(n))
        c = hoja.cell(row=fila_datos, column=col, value=est)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor={"OK": "2E8B57", "CON PROBLEMAS": "C0392B",
                                                "SIN MEDIR": "7F8C8D"}[est])
        col += 1
    for cc in range(1, total_cols + 1):
        celda = hoja.cell(row=fila_datos, column=cc)
        celda.border = borde
        celda.alignment = Alignment(horizontal="center", vertical="center")
        if cc == 1:
            celda.font = Font(bold=True)
            celda.fill = PatternFill("solid", fgColor="E8EEF4")
        elif fila_datos % 2 == 1 and celda.fill.fgColor.rgb in ("00000000", None):
            celda.fill = PatternFill("solid", fgColor="F3F6F9")
    hoja.row_dimensions[fila_datos].height = 26
    fila_datos += 1
ultima = fila_datos - 1

# Barras de datos en ingesta y entrega
for letra in ("D", "E", "F", "G"):
    hoja.conditional_formatting.add(
        f"{letra}6:{letra}{ultima}",
        DataBarRule(start_type="num", start_value=0, end_type="max", color="F5B041"))

# Anchos
hoja.column_dimensions["A"].width = 12
for i in range(2, total_cols + 1):
    hoja.column_dimensions[get_column_letter(i)].width = 14

# Leyenda
r = ultima + 2
leyenda = [
    "Cómo leer esta tabla:",
    "• Concurrencia: clientes que lograron conectarse a la vez. Si es menor que los clientes solicitados, el servidor rechazó conexiones.",
    "• Ingesta: alertas que el servidor aceptó por segundo (validadas, guardadas en SQLite y difundidas).",
    "• Entrega: mensajes recibidos por los clientes suscritos, por segundo.",
    "• Estado OK = sin conexiones fallidas, sin errores, sin desconexiones y con todos los mensajes entregados.",
    "• n/d = no se pudo medir (el servidor rechazó nuevas conexiones antes de iniciar la carga).",
]
for i, texto in enumerate(leyenda):
    hoja.merge_cells(start_row=r + i, start_column=1, end_row=r + i, end_column=total_cols)
    c = hoja.cell(row=r + i, column=1, value=texto)
    c.font = Font(bold=(i == 0), size=10, color="444444")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

# Graficos (datos auxiliares en hoja "Graficos")
aux = libro.create_sheet("Graficos")
aux.append(["Clientes", "Flask", "Node.js"])
for n in escenarios:
    fila = [str(n)]
    for srv in ("Flask", "Node.js"):
        r_ = datos[srv].get(n)
        v = num(r_["throughput_ingesta_alertas_s"]) if r_ and int(r_["emisores"]) > 0 else 0
        fila.append(0 if v == "n/d" else v)
    aux.append(fila)
aux.append([])
aux.append(["Clientes", "Flask", "Node.js"])
ini_conc = aux.max_row
for n in escenarios:
    fila = [str(n)]
    for srv in ("Flask", "Node.js"):
        r_ = datos[srv].get(n)
        fila.append(num(r_["conexiones_simultaneas"]) if r_ else 0)
    aux.append(fila)


def grafico(titulo, eje_y, hdr, primera, ultima_f, ancla):
    ch = BarChart()
    ch.type = "col"
    ch.title = titulo
    ch.y_axis.title = eje_y
    ch.x_axis.title = "Clientes solicitados"
    ch.height, ch.width = 8.5, 16
    datos_ref = Reference(aux, min_col=2, max_col=3, min_row=hdr, max_row=ultima_f)
    cats = Reference(aux, min_col=1, min_row=primera, max_row=ultima_f)
    ch.add_data(datos_ref, titles_from_data=True)
    ch.set_categories(cats)
    ch.series[0].graphicalProperties.solidFill = COLOR_FLASK
    ch.series[1].graphicalProperties.solidFill = COLOR_NODE
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    hoja.add_chart(ch, ancla)


fila_graf = ultima + 2 + len(leyenda) + 1
grafico("Alertas aceptadas por segundo", "alertas/s", 1, 2, 1 + len(escenarios), f"A{fila_graf}")
grafico("Conexiones simultáneas logradas", "clientes conectados",
        ini_conc, ini_conc + 1, ini_conc + len(escenarios), f"H{fila_graf}")

hoja.freeze_panes = "B6"
hoja.page_setup.orientation = "landscape"
hoja.sheet_properties.pageSetUpPr.fitToPage = True
hoja.page_setup.fitToWidth = 1
hoja.page_setup.fitToHeight = 0

# Hoja de detalle (formato largo)
det = libro.create_sheet("Detalle", 1)
det.append(columnas)
for f in filas:
    det.append([num(x) if i >= 2 and x not in ("", "no medido") else x for i, x in enumerate(f)])
for celda in det[1]:
    celda.font = Font(bold=True, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor=AZUL)
    celda.alignment = centro
for fila in det.iter_rows(min_row=2):
    for celda in fila:
        celda.alignment = Alignment(horizontal="center")
        celda.border = borde
for columna in det.columns:
    det.column_dimensions[get_column_letter(columna[0].column)].width = 16
det.row_dimensions[1].height = 30
det.freeze_panes = "A2"
det.auto_filter.ref = det.dimensions

archivo_xlsx = base / "tabla_comparativa_throughput.xlsx"
libro.save(archivo_xlsx)

print(f"\nCSV guardado en:\n{archivo_csv}")
print(f"Excel guardado en:\n{archivo_xlsx}")

try:
    os.startfile(archivo_xlsx)   # abre el Excel directo (Windows)
except Exception:
    pass