import asyncio
import csv
import os
import sys
import time
from datetime import datetime
from pathlib import Path

import socketio
from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

URL = sys.argv[1] if len(sys.argv) > 1 else "http://localhost:5000"
NOMBRE = sys.argv[2] if len(sys.argv) > 2 else ("flask" if ":5000" in URL else "node")

if len(sys.argv) > 3 and sys.argv[3]:
    ESCENARIOS = [int(x) for x in sys.argv[3].split(",")]
    SUFIJO = "_" + "-".join(str(x) for x in ESCENARIOS)
else:
    ESCENARIOS = [50, 100, 500, 1000]
    SUFIJO = ""

DURACION = float(sys.argv[4]) if len(sys.argv) > 4 else 10.0   # segundos de carga
EMISORES = 5               # emisores concurrentes (lazo cerrado)
RECEPTORES = 10            # clientes suscritos a la zona destino
ZONA_DESTINO = "norte"
ZONAS_RESTO = ["centro", "sur"]
TIMEOUT_CONEXION = 8
TIMEOUT_ACK = 5
INTERVALO_CREACION = 0.005
ESPERA_SUSCRIPCION = 2
ESPERA_FINAL = 3

resultados = []
ARCHIVO = Path(__file__).with_name(f"resultados_throughput_{NOMBRE}{SUFIJO}.csv")


def registrar(fila):
    # Guarda el CSV completo despues de cada escenario.
    resultados.append(fila)
    print(fila)
    with ARCHIVO.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=resultados[0].keys())
        w.writeheader()
        w.writerows(resultados)


def guardar_excel():
    archivo_excel = ARCHIVO.with_suffix(".xlsx")
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Resultados"
    hoja.sheet_view.showGridLines = False

    AZUL = "245A81"
    fino = Side(style="thin", color="BFC9D4")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    nombre_visible = "Flask" if NOMBRE == "flask" else ("Node.js" if NOMBRE == "node" else NOMBRE)

    # (grupo, titulo, clave, color del grupo, formato)
    cols = [
        ("CARGA", "Clientes\nsolicitados", "clientes_solicitados", "245A81", "0"),
        ("CARGA", "Conexiones\nsimultáneas", "conexiones_simultaneas", "245A81", "0"),
        ("CARGA", "Conexiones\nfallidas", "conexiones_fallidas", "245A81", "0"),
        ("MENSAJES", "Alertas\nenviadas", "alertas_enviadas", "2E7D6B", "#,##0"),
        ("MENSAJES", "Alertas\naceptadas", "alertas_aceptadas", "2E7D6B", "#,##0"),
        ("MENSAJES", "Mensajes\nesperados", "mensajes_esperados", "2E7D6B", "#,##0"),
        ("MENSAJES", "Mensajes\nrecibidos", "mensajes_recibidos", "2E7D6B", "#,##0"),
        ("RENDIMIENTO", "Ingesta\n(alertas/s)", "throughput_ingesta_alertas_s", "B26A00", "#,##0.00"),
        ("RENDIMIENTO", "Entrega\n(mensajes/s)", "throughput_entrega_msgs_s", "B26A00", "#,##0.00"),
        ("PROBLEMAS", "Errores", None, "8E3B46", "0"),
        ("PROBLEMAS", "Desconexiones\ninesperadas", "desconexiones_inesperadas", "8E3B46", "0"),
        ("RESULTADO", "Estado", None, "4A4A4A", "@"),
        ("RESULTADO", "Observación", "nota", "4A4A4A", "@"),
    ]
    n = len(cols)

    # Titulo y datos de la prueba
    hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    hoja["A1"] = f"Throughput y concurrencia — {nombre_visible}"
    hoja["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    hoja["A1"].fill = PatternFill("solid", fgColor=AZUL)
    hoja["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    hoja.row_dimensions[1].height = 32

    primera = resultados[0]
    info = (f"Servidor: {primera['url']}   |   Fecha: {primera['fecha']}   |   "
            f"Carga por escenario: {DURACION:g} s   |   "
            f"Emisores simultáneos: {EMISORES}")
    hoja.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    hoja["A2"] = info
    hoja["A2"].font = Font(italic=True, size=10, color="555555")
    hoja["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    hoja.row_dimensions[2].height = 20

    # Fila de grupos (4) y de titulos (5)
    inicio = 1
    while inicio <= n:
        grupo = cols[inicio - 1][0]
        fin_g = inicio
        while fin_g < n and cols[fin_g][0] == grupo:
            fin_g += 1
        hoja.merge_cells(start_row=4, start_column=inicio, end_row=4, end_column=fin_g)
        celda = hoja.cell(row=4, column=inicio, value=grupo)
        for c in range(inicio, fin_g + 1):
            hoja.cell(row=4, column=c).fill = PatternFill("solid", fgColor=cols[inicio - 1][3])
            hoja.cell(row=4, column=c).border = borde
        celda.font = Font(bold=True, color="FFFFFF", size=10)
        celda.alignment = centro
        inicio = fin_g + 1
    for i, (grupo, titulo, clave, color, fmt) in enumerate(cols, start=1):
        c = hoja.cell(row=5, column=i, value=titulo)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=color)
        c.alignment = centro
        c.border = borde
    hoja.row_dimensions[4].height = 20
    hoja.row_dimensions[5].height = 34

    # Datos
    for r, fila in enumerate(resultados, start=6):
        errores = int(fila["errores_rechazo"]) + int(fila["errores_timeout"])
        sin_medir = fila["emisores"] == 0
        ok = (not sin_medir and fila["conexiones_fallidas"] == 0 and errores == 0
              and fila["desconexiones_inesperadas"] == 0
              and fila["mensajes_recibidos"] == fila["mensajes_esperados"])
        estado = "SIN MEDIR" if sin_medir else ("OK" if ok else "CON PROBLEMAS")
        for i, (grupo, titulo, clave, color, fmt) in enumerate(cols, start=1):
            if clave:
                valor = fila.get(clave, "")
            elif titulo == "Errores":
                valor = errores
            else:
                valor = estado
            if valor == "" and fmt != "@":
                valor = "n/d"
            c = hoja.cell(row=r, column=i, value=valor)
            c.number_format = fmt
            c.border = borde
            c.alignment = Alignment(horizontal="left" if clave == "nota" else "center",
                                    vertical="center", wrap_text=(clave == "nota"))
            if (r % 2) == 1:
                c.fill = PatternFill("solid", fgColor="F3F6F9")
        est = hoja.cell(row=r, column=n - 1)
        est.font = Font(bold=True, color="FFFFFF")
        est.fill = PatternFill("solid", fgColor={"OK": "2E8B57", "CON PROBLEMAS": "C0392B",
                                                  "SIN MEDIR": "7F8C8D"}[estado])
        hoja.row_dimensions[r].height = 24

    ultima = 5 + len(resultados)
    # Barras de datos en los dos throughput
    for col in ("H", "I"):
        hoja.conditional_formatting.add(
            f"{col}6:{col}{ultima}",
            DataBarRule(start_type="num", start_value=0, end_type="max", color="F5B041"))

    # Anchos
    anchos = [12, 14, 12, 12, 12, 12, 12, 14, 15, 10, 15, 16, 48]
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho

    # Leyenda
    r = ultima + 2
    leyenda = [
        "Cómo leer esta tabla:",
        "• Ingesta (alertas/s): alertas que el servidor aceptó por segundo (ya validadas, guardadas en SQLite y difundidas).",
        "• Entrega (mensajes/s): mensajes recibidos por los clientes suscritos, por segundo.",
        "• Conexiones simultáneas: clientes que lograron conectarse a la vez durante la prueba.",
        "• Estado OK = sin conexiones fallidas, sin errores, sin desconexiones y con todos los mensajes entregados.",
    ]
    for i, texto in enumerate(leyenda):
        hoja.merge_cells(start_row=r + i, start_column=1, end_row=r + i, end_column=n)
        c = hoja.cell(row=r + i, column=1, value=texto)
        c.font = Font(bold=(i == 0), size=10, color="444444")
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    hoja.freeze_panes = "A6"
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_setup.fitToHeight = 0
    libro.save(archivo_excel)
    print(f"Excel guardado en:\n{archivo_excel}")
    try:
        os.startfile(archivo_excel)   # abre el Excel directo (Windows)
    except Exception:
        pass


async def cerrar(cliente):
    try:
        if cliente.connected:
            await cliente.disconnect()
        sesion = getattr(cliente.eio, "http", None)
        if sesion is not None and not sesion.closed:
            await sesion.close()
    except Exception:
        pass


async def medir_escenario(cantidad):
    print(f"\n--- Escenario: {cantidad} clientes ({NOMBRE}) ---")

    c = {
        "recibidos": 0,
        "enviadas": 0,
        "aceptadas": 0,
        "err_rechazo": 0,
        "err_timeout": 0,
        "inesperadas": 0,
        "cierre_intencional": False,
    }
    clientes = []
    n_receptores = min(RECEPTORES, cantidad)

    async def crear_cliente(indice):
        zona = ZONA_DESTINO if indice < n_receptores else ZONAS_RESTO[indice % 2]
        cli = socketio.AsyncClient(reconnection=False)

        if zona == ZONA_DESTINO:
            @cli.on("alerta_recibida")
            async def _on_alerta(data):
                c["recibidos"] += 1

        @cli.event
        async def disconnect():
            if not c["cierre_intencional"]:
                c["inesperadas"] += 1

        try:
            await asyncio.wait_for(
                cli.connect(URL, transports=["websocket"]),
                timeout=TIMEOUT_CONEXION,
            )
            await cli.emit("suscribir_zona", {"zona": zona})
            clientes.append(cli)
        except Exception:
            await cerrar(cli)

    # 1) Conectar gradualmente (misma tecnica que prueba_carga.py).
    tareas = []
    for i in range(cantidad):
        tareas.append(asyncio.create_task(crear_cliente(i)))
        await asyncio.sleep(INTERVALO_CREACION)
    await asyncio.gather(*tareas)

    conectados = len(clientes)
    fallidas = cantidad - conectados
    print(f"Conectados: {conectados}/{cantidad}")
    await asyncio.sleep(ESPERA_SUSCRIPCION)

    # 2) Emisores en lazo cerrado.
    emisores = []
    for _ in range(EMISORES):
        e = socketio.AsyncClient(reconnection=False)
        try:
            await e.connect(URL, transports=["websocket"])
            emisores.append(e)
        except Exception as error:
            print(f"Un emisor no pudo conectarse: {error}")
            await cerrar(e)

    if not emisores:
        await asyncio.gather(*(cerrar(x) for x in clientes))
        registrar(armar_fila(cantidad, conectados, fallidas, c, 0.0, 0,
                             "el servidor rechazo nuevas conexiones; no se pudo medir"))
        return

    fin = time.perf_counter() + DURACION

    async def lazo(e, id_emisor):
        n = 0
        while time.perf_counter() < fin:
            n += 1
            c["enviadas"] += 1
            alerta = {"tipo": "incendio", "gravedad": "alta", "zona": ZONA_DESTINO,
                      "mensaje": f"thr-{id_emisor}-{n}"}
            try:
                resp = await e.call("emitir_alerta", alerta, timeout=TIMEOUT_ACK)
                if isinstance(resp, dict) and resp.get("ok"):
                    c["aceptadas"] += 1
                else:
                    c["err_rechazo"] += 1
            except Exception:
                c["err_timeout"] += 1

    t0 = time.perf_counter()
    await asyncio.gather(*(lazo(e, i) for i, e in enumerate(emisores)))
    duracion_real = time.perf_counter() - t0

    # Tiempo para que lleguen los ultimos mensajes en transito.
    await asyncio.sleep(ESPERA_FINAL)

    c["cierre_intencional"] = True
    await asyncio.gather(*(cerrar(x) for x in clientes + emisores))

    nota = ""
    if fallidas:
        nota = f"solo se conectaron {conectados} de {cantidad} clientes"
    registrar(armar_fila(cantidad, conectados, fallidas, c, duracion_real,
                         len(emisores), nota))


def armar_fila(cantidad, conectados, fallidas, c, duracion, n_emisores, nota=""):
    esperados = c["aceptadas"] * min(RECEPTORES, conectados)
    return {
        "fecha": datetime.now().isoformat(timespec="seconds"),
        "servidor": NOMBRE,
        "url": URL,
        "clientes_solicitados": cantidad,
        "conexiones_simultaneas": conectados,
        "conexiones_fallidas": fallidas,
        "emisores": n_emisores,
        "duracion_carga_s": round(duracion, 2),
        "alertas_enviadas": c["enviadas"],
        "alertas_aceptadas": c["aceptadas"],
        "errores_rechazo": c["err_rechazo"],
        "errores_timeout": c["err_timeout"],
        "mensajes_esperados": esperados,
        "mensajes_recibidos": c["recibidos"],
        "desconexiones_inesperadas": c["inesperadas"],
        "throughput_ingesta_alertas_s": round(c["aceptadas"] / duracion, 2) if duracion else "",
        "throughput_entrega_msgs_s": round(c["recibidos"] / duracion, 2) if duracion else "",
        "nota": nota,
    }


async def main():
    print(f"Servidor: {URL} ({NOMBRE}) | duracion de carga: {DURACION}s | emisores: {EMISORES}")
    for cantidad in ESCENARIOS:
        await medir_escenario(cantidad)
        await asyncio.sleep(1.5)
    print(f"\nListo. Resultados guardados en:\n{ARCHIVO}")
    guardar_excel()


if __name__ == "__main__":
    asyncio.run(main())