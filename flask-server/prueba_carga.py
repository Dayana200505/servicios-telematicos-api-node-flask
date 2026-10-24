
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import asyncio
import csv
import sys 
import time
from datetime import datetime
from pathlib import Path

import socketio

URL = sys.argv[1] if len(sys.argv) > 1 else "http://localhost:5000"

ESCENARIOS = [50, 100, 500, 1000]
TIMEOUT_CONEXION = 8
DURACION_CONEXION = 5
INTERVALO_CREACION = 0.005

resultados = []


async def probar_escenario(cantidad):
    
    print(f"\n--- Escenario: {cantidad} clientes ---")
    print(f"Servidor: {URL}")

    exitosas = 0
    fallidas = 0
    inesperadas = 0
    controladas = 0
    tiempos = []
    clientes = []
    inicio = time.perf_counter()

    async def crear_cliente(indice):
        nonlocal exitosas, fallidas

        cliente = socketio.AsyncClient(
            reconnection=False,
            logger=False,
            engineio_logger=False,
        )

        estado = {
            "conectado": False,
            "contado": False,
            "cierre_intencional": False,
        }

        inicio_conexion = time.perf_counter()

        @cliente.event
        async def disconnect():
            nonlocal inesperadas, controladas

            if not estado["conectado"]:
                return

            estado["conectado"] = False

            if estado["cierre_intencional"]:
                controladas += 1
            else:
                inesperadas += 1

        try:
            await asyncio.wait_for(
                cliente.connect(URL, transports=["websocket"]),
                timeout=TIMEOUT_CONEXION,
            )

            estado["conectado"] = True
            estado["contado"] = True
            exitosas += 1
            tiempos.append(
                (time.perf_counter() - inicio_conexion) * 1000
            )

        except Exception as error:
            if not estado["contado"]:
                estado["contado"] = True
                fallidas += 1

                if indice < 3:
                    print(f"Cliente {indice + 1}: fallo: {error}")

        clientes.append((cliente, estado))

    # Crear conexiones de forma gradual.
    tareas = []

    for indice in range(cantidad):
        tareas.append(asyncio.create_task(crear_cliente(indice)))
        await asyncio.sleep(INTERVALO_CREACION)

    await asyncio.gather(*tareas)

    print(f"Clientes solicitados: {cantidad}")
    print(f"Conexiones exitosas: {exitosas}")
    print(f"Conexiones fallidas: {fallidas}")

    promedio = sum(tiempos) / len(tiempos) if tiempos else None
    print(
        "Tiempo promedio de conexión: "
        + (f"{promedio:.2f} ms" if promedio is not None else "N/A")
    )

    print(f"Manteniendo conexiones durante {DURACION_CONEXION} segundos...")
    await asyncio.sleep(DURACION_CONEXION)

    # Cerrar los clientes de forma intencional.
    for cliente, estado in clientes:
        if cliente.connected:
            estado["cierre_intencional"] = True
            await cliente.disconnect()
        
        # Cerrar explícitamente la sesión HTTP interna, si existe.
        sesion = getattr(cliente.eio, "http", None)

        if sesion is not None and not sesion.closed:
            await sesion.close()

    await asyncio.sleep(0.2)

    resultado = {
        "fecha": datetime.now().isoformat(timespec="seconds"),
        "servidor": URL,
        "clientes_solicitados": cantidad,
        "conexiones_exitosas": exitosas,
        "conexiones_fallidas": fallidas,
        "desconexiones_inesperadas": inesperadas,
        "desconexiones_controladas": controladas,
        "tiempo_promedio_conexion_ms": (
            round(promedio, 2) if promedio is not None else ""
        ),
        "duracion_total_ms": round(
            (time.perf_counter() - inicio) * 1000, 2
        ),
    }

    resultados.append(resultado)
    print(resultado)


async def main():
    for cantidad in ESCENARIOS:
        await probar_escenario(cantidad)
        await asyncio.sleep(1.5)

    archivo = Path(__file__).with_name("resultados_carga_flask.csv")

    with archivo.open("w", newline="", encoding="utf-8") as f:
        escritor = csv.DictWriter(
            f,
            fieldnames=resultados[0].keys(),
        )
        escritor.writeheader()
        escritor.writerows(resultados)

    print(f"\nPrueba finalizada. Resultados guardados en:\n{archivo}")


    # Crear también un Excel organizado para revisar los resultados.
    archivo_excel = Path(__file__).with_name(
        "resultados_carga_flask.xlsx"
    )

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Resultados Flask"

    columnas = list(resultados[0].keys())
    hoja.append(columnas)

    for resultado in resultados:
        hoja.append([resultado.get(columna, "") for columna in columnas])

  
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(
            fill_type="solid",
            fgColor="245A81",
        )
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    # Ajustar anchos y activar filtros.
    for columna in hoja.columns:
        letra = get_column_letter(columna[0].column)
        ancho = max(
            len(str(celda.value or "")) for celda in columna
        )
        hoja.column_dimensions[letra].width = min(ancho + 3, 32)

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    libro.save(archivo_excel)

    print(f"Excel organizado guardado en:\n{archivo_excel}")


if __name__ == "__main__":
    asyncio.run(main())
